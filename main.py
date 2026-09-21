import os
import shutil
from fastapi import FastAPI, File, UploadFile, Form
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from crewai import Crew, Process
import docx

# Import your agents and tasks from your other files
# (Adjust these import names if your files/variables are named differently)
from agents import ats_expert, resume_writer, cover_letter_crafter
from tasks import analyze_resume_task, rewrite_resume_task, write_cover_letter_task

app = FastAPI()

# --- CORS CONFIGURATION ---
# This allows your Vercel frontend to talk to this Render backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://atsresume-eight.vercel.app", 
        "http://localhost:5173" # For local testing
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- MAIN ENDPOINT ---
# Using 'def' instead of 'async def' so CrewAI runs in a background thread
@app.post("/api/optimize")
def optimize_resume(
    job_description: str = Form(...),
    resume: UploadFile = File(...)
):
    print(f"DEBUG: Request received! File: {resume.filename}")
    
    temp_resume_path = f"temp_{resume.filename}"
    with open(temp_resume_path, "wb") as buffer:
        shutil.copyfileobj(resume.file, buffer)
        
    try:
        # EXTRACT TEXT FROM THE DOCX FILE
        print("DEBUG: Extracting text from document...")
        doc = docx.Document(temp_resume_path)
        extracted_resume_text = "\n".join([para.text for para in doc.paragraphs])

        print("DEBUG: Initializing CrewAI...")
        my_crew = Crew(
            agents=[ats_expert, resume_writer, cover_letter_crafter],
            tasks=[analyze_resume_task, rewrite_resume_task, write_cover_letter_task],
            process=Process.sequential,
            verbose=True
        )

        print("DEBUG: Starting CrewAI kickoff...")
        # UPDATE THIS DICTIONARY TO MATCH YOUR TASKS
       # 3. Run the AI process
        print("DEBUG: Starting CrewAI kickoff...")
        my_crew.kickoff(inputs={
            'resume_text': extracted_resume_text,
            'job_description': job_description
        })
        
        # 4. BUILD AND SAVE THE WORD DOCUMENT
        print("DEBUG: Building the final Word document...")
        output_doc = docx.Document()
        
        # Add the Cover Letter (from Task 3)
        output_doc.add_heading('Cover Letter', level=1)
        output_doc.add_paragraph(str(write_cover_letter_task.output))
        
        output_doc.add_page_break()
        
        # Add the Optimized Resume (from Task 2)
        output_doc.add_heading('Optimized Resume', level=1)
        output_doc.add_paragraph(str(rewrite_resume_task.output))
        
        # Save it to the server's hard drive
        output_filename = "Optimized_Application.docx"
        output_doc.save(output_filename)
        
        # 5. Return the generated document
        print("DEBUG: Process complete, sending file back to frontend!")
        return FileResponse(
            path=output_filename, 
            filename=output_filename,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        
    except Exception as e:
        # Catch and print any hidden crashes to the Render logs
        print(f"CRASH REPORT: {str(e)}") 
        raise e
        
    finally:
        # 5. Clean up the temporary uploaded resume so your server storage doesn't fill up
        if os.path.exists(temp_resume_path):
            os.remove(temp_resume_path)
            print(f"DEBUG: Cleaned up temporary file {temp_resume_path}")